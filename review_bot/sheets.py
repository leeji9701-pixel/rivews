"""확인대기.xlsx (사람이 답글 작성) 와 처리내역/날짜.xlsx (매일 결과) 관리."""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

log = logging.getLogger(__name__)

PENDING_HEADERS = [
    "상태", "등록일", "플랫폼", "계정", "고객", "주문횟수", "별점",
    "리뷰내용", "분류사유", "캡처파일", "답글(여기에 작성)", "리뷰키(수정금지)",
]
REPORT_HEADERS = [
    "시간", "플랫폼", "계정", "고객", "주문횟수", "별점", "사진",
    "리뷰내용", "분류", "사유", "템플릿", "답글", "결과",
]
WAITING, DONE = "대기", "등록완료"
COL = {h: i + 1 for i, h in enumerate(PENDING_HEADERS)}


def _new_book(headers: list[str], widths: dict[int, int]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="8B4513")
    for col, w in widths.items():
        ws.column_dimensions[chr(64 + col)].width = w
    ws.freeze_panes = "A2"
    return wb


def _save(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
        return path
    except PermissionError:  # 엑셀에서 파일을 열어둔 경우
        alt = path.with_name(f"{path.stem}_{datetime.now():%H%M%S}{path.suffix}")
        wb.save(alt)
        log.warning("%s 파일이 열려 있어 %s 로 저장했습니다.", path.name, alt.name)
        return alt


class PendingSheet:
    def __init__(self, path: Path):
        self.path = path
        if path.exists():
            self.wb = load_workbook(path)
        else:
            self.wb = _new_book(PENDING_HEADERS, {1: 9, 2: 11, 3: 10, 4: 12, 5: 12, 8: 50, 9: 22, 10: 40, 11: 60, 12: 18})
        self.ws = self.wb.active

    def keys(self) -> set[str]:
        return {str(r[COL["리뷰키(수정금지)"] - 1]) for r in self.ws.iter_rows(min_row=2, values_only=True) if r[0]}

    def add(self, review, reason: str, capture: str) -> None:
        self.ws.append([
            WAITING, datetime.now().strftime("%Y-%m-%d"), review.platform, review.account,
            review.author, review.order_count, review.rating, review.text, reason, capture, "", review.key,
        ])
        for c in self.ws[self.ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    def approved(self, platform: str, account: str) -> list[tuple[int, str, str]]:
        """(행번호, 리뷰키, 답글) - 상태가 '대기'이고 답글이 채워진 것."""
        out = []
        for i, r in enumerate(self.ws.iter_rows(min_row=2, values_only=True), start=2):
            status, plat, acc = r[0], r[COL["플랫폼"] - 1], r[COL["계정"] - 1]
            reply = (r[COL["답글(여기에 작성)"] - 1] or "").strip()
            if status == WAITING and plat == platform and acc == account and reply:
                out.append((i, str(r[COL["리뷰키(수정금지)"] - 1]), reply))
        return out

    def mark_done(self, row: int) -> None:
        self.ws.cell(row=row, column=COL["상태"], value=DONE)

    def save(self) -> Path:
        return _save(self.wb, self.path)


class DailyReport:
    def __init__(self, folder: Path):
        self.path = folder / f"{datetime.now():%Y-%m-%d}.xlsx"
        self.rows: list[list] = []

    def add(self, review, kind: str, reason: str, template: str, reply: str, result: str) -> None:
        self.rows.append([
            datetime.now().strftime("%H:%M"), review.platform, review.account, review.author,
            review.order_count, review.rating, "O" if review.has_photo else "", review.text,
            kind, reason, template, reply, result,
        ])

    def save(self) -> Path | None:
        if not self.rows:
            return None
        if self.path.exists():
            wb = load_workbook(self.path)
        else:
            wb = _new_book(REPORT_HEADERS, {4: 12, 8: 50, 10: 22, 12: 60, 13: 14})
        ws = wb.active
        for row in self.rows:
            ws.append(row)
            for c in ws[ws.max_row]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        return _save(wb, self.path)
