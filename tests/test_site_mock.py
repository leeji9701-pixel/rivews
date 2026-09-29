"""가짜 리뷰 페이지로 로그인→분류→답글→캡처 흐름 전체를 검증."""
import os
from pathlib import Path

import pytest

pytest.importorskip("playwright")
from playwright.sync_api import sync_playwright  # noqa: E402

from review_bot.config import Account  # noqa: E402
from review_bot.site import ReviewSite  # noqa: E402

HTML = """
<div class="list">
 <div class="card"><b class="nick">김철수</b><span class="ord">7번째 주문</span>
  <i class="star on"></i><i class="star on"></i><i class="star on"></i><i class="star on"></i><i class="star on"></i>
  <p class="txt">바삭하고 맛있어요</p><img class="pic" src="x.png">
  <button class="rb" onclick="this.parentNode.querySelector('.form').style.display='block'">댓글</button>
  <div class="form" style="display:none"><textarea></textarea>
   <button onclick="const c=this.closest('.card');const d=document.createElement('div');d.className='reply';d.textContent=c.querySelector('textarea').value;c.appendChild(d);this.parentNode.remove()">등록</button></div>
 </div>
 <div class="card"><b class="nick">이영희</b><i class="star on"></i><p class="txt">최악이에요</p><button class="rb">댓글</button></div>
 <div class="card"><b class="nick">박민수</b><i class="star on"></i><i class="star on"></i><i class="star on"></i><i class="star on"></i>
  <p class="txt">좋아요</p><div class="reply">이미 답글</div></div>
</div>"""

SEL = {
    "reviews_url": "about:blank",
    "login": {"url": "", "id_input": "#id", "pw_input": "#pw", "submit": "#go"},
    "reviews": {
        "list_ready": ".list", "unanswered_tab": "", "card": ".card", "author": ".nick",
        "order_count_text": ".ord", "order_count_regex": r"(\d+)\s*번째", "date": "",
        "rating": {"mode": "count", "selector": ".star.on"}, "text": ".txt", "photo": "img.pic",
        "reply_exists": ".reply", "reply_button": ".rb", "reply_textarea": "textarea",
        "reply_submit": "button:has-text('등록')", "load_more": "",
    },
}


@pytest.fixture
def site():
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=os.environ.get("CHROMIUM_PATH") or None)
        page = browser.new_page()
        page.set_content(HTML)
        yield ReviewSite(page, Account("baemin", "테스트", "T"), SEL, delay=(0, 0))
        browser.close()


def test_extract_skips_answered(site):
    reviews = site.unanswered()
    assert [x.author for x in reviews] == ["김철수", "이영희"]
    kim = reviews[0]
    assert (kim.rating, kim.order_count, kim.has_photo) == (5, 7, True)


def test_post_reply_and_capture(site, tmp_path: Path):
    kim, lee = site.unanswered()
    assert site.post_reply(kim.key, "감사합니다")
    assert [x.author for x in site.unanswered()] == ["이영희"]
    shot = tmp_path / "lee.png"
    assert site.capture(lee.key, shot) and shot.stat().st_size > 0


def test_runner_never_posts_manual_reviews(tmp_path, monkeypatch):
    """전체 흐름: 5점은 자동 등록, 1점(최악)은 등록 없이 캡처 + 확인대기 초안만."""
    from openpyxl import load_workbook

    from review_bot import runner
    from review_bot.config import Settings

    settings = Settings({
        "dry_run": False, "headless": True, "delay_seconds": [0, 0], "max_replies_per_account": 50,
        "folders": {"captures": str(tmp_path / "cap"), "reports": str(tmp_path / "rep"),
                    "pending_file": str(tmp_path / "pending.xlsx")},
        "accounts": [{"platform": "baemin", "name": "테스트", "env_prefix": "T"}],
    })
    real_yaml = runner.load_yaml
    monkeypatch.setattr(runner, "load_yaml", lambda n: {"baemin": SEL} if n == "selectors.yaml" else real_yaml(n))
    monkeypatch.setattr(runner, "DATA_DIR", tmp_path)

    pages = []

    def fake_ctx(p, account, headless, channel=None):
        ctx = p.chromium.launch_persistent_context(
            str(tmp_path / "s"), headless=True, executable_path=os.environ.get("CHROMIUM_PATH") or None)
        ctx.pages[0].set_content(HTML)
        pages.append(ctx.pages[0])
        ctx.pages[0].goto = lambda *a, **k: None  # 가짜 페이지 유지
        return ctx

    monkeypatch.setattr(runner, "_open_context", fake_ctx)

    runner.run(settings)

    ws = load_workbook(tmp_path / "pending.xlsx").active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    status, *_, capture, draft, _key = rows[0]
    assert status == "검수대기" and rows[0][4] == "이영희"
    assert draft.startswith("101번지 남산돈까스 본점입니다.")
    assert list((tmp_path / "cap").rglob("*.png"))

    report = load_workbook(next((tmp_path / "rep").glob("*.xlsx"))).active
    results = {r[3]: r[12] for r in report.iter_rows(min_row=2, values_only=True)}
    assert results == {"김철수": "등록완료", "이영희": "검수대기"}
